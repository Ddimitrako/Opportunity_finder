from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook
from pydantic import ValidationError

from app.catalog import SoftwareCatalog, catalog_version


DELIMITER = " | "
EXPECTED_COUNTS = {
    "products": 63,
    "departments": 8,
    "categories": 39,
    "licenses": 11,
    "repository_health": 63,
}
BACKEND_DIR = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = BACKEND_DIR / "catalog" / "software_catalog.xlsx"
DEFAULT_OUTPUT = Path(__file__).resolve().parent / "data" / "software_catalog.json"


def _rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing worksheet: {sheet_name}")
    worksheet = workbook[sheet_name]
    values = list(worksheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"Worksheet is empty: {sheet_name}")
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    if not all(headers):
        raise ValueError(f"Worksheet has blank headers: {sheet_name}")
    output: list[dict[str, Any]] = []
    for raw in values[1:]:
        if not any(value not in (None, "") for value in raw):
            continue
        output.append({headers[index]: raw[index] for index in range(len(headers))})
    return output


def _text(value: Any, *, required: bool = False) -> str | None:
    if value is None:
        if required:
            raise ValueError("Required text value is missing")
        return None
    output = str(value).strip()
    if not output and required:
        raise ValueError("Required text value is blank")
    return output or None


def _list(value: Any) -> list[str]:
    text = _text(value)
    if not text:
        return []
    return [item.strip() for item in text.split(DELIMITER) if item.strip()]


def _bool(value: Any, *, default: bool = False) -> bool:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"true", "yes", "1"}:
        return True
    if normalized in {"false", "no", "0"}:
        return False
    raise ValueError(f"Invalid boolean value: {value}")


def _date(value: Any) -> str:
    if hasattr(value, "date"):
        return value.date().isoformat()
    text = _text(value, required=True)
    assert text is not None
    return text[:10]


def _iso(value: Any) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat().replace("+00:00", "Z")
    return _text(value)


def workbook_payload(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    dictionary = {
        str(row["key"]).strip(): str(row["value"]).strip()
        for row in _rows(workbook, "Data Dictionary")
        if row.get("key") not in (None, "")
    }
    schema_version = int(dictionary.get("schema_version", "1"))

    departments = [
        {
            "id": _text(row["id"], required=True),
            "name": _text(row["name"], required=True),
            "short_name": _text(row["short_name"], required=True),
            "description": _text(row["description"], required=True),
            "icon": _text(row["icon"], required=True),
        }
        for row in _rows(workbook, "Departments")
    ]
    categories = [
        {
            "id": _text(row["id"], required=True),
            "department_id": _text(row["department_id"], required=True),
            "name": _text(row["name"], required=True),
            "description": _text(row["description"], required=True),
            "cpv_prefixes": _list(row.get("cpv_prefixes")),
            "keywords_en": _list(row.get("keywords_en")),
            "keywords_el": _list(row.get("keywords_el")),
            "negative_keywords": _list(row.get("negative_keywords")),
        }
        for row in _rows(workbook, "Categories")
    ]
    licenses = [
        {
            "id": _text(row["id"], required=True),
            "spdx": _text(row["spdx"], required=True),
            "name": _text(row["name"], required=True),
            "family": _text(row["family"], required=True),
            "osi_approved": _bool(row.get("osi_approved"), default=True),
            "summary": _text(row["summary"], required=True),
            "obligations": _list(row.get("obligations")),
            "evidence": _text(row["evidence"], required=True),
        }
        for row in _rows(workbook, "Licenses")
    ]
    repository_health = [
        {
            "slug": _text(row["slug"], required=True),
            "repo": _text(row["repo"], required=True),
            "stars": int(row["stars"]) if row.get("stars") not in (None, "") else None,
            "latest_release": _text(row.get("latest_release")),
            "last_activity": _iso(row.get("last_activity")),
            "archived": _bool(row.get("archived")),
            "observed_license": _text(row.get("observed_license")),
            "license_drift": _bool(row.get("license_drift")),
            "verified_at": _iso(row.get("verified_at")),
            "status": _text(row["status"], required=True),
        }
        for row in _rows(workbook, "Repository Health")
    ]
    products = [
        {
            "slug": _text(row["slug"], required=True),
            "name": _text(row["name"], required=True),
            "edition": _text(row["edition"], required=True),
            "repository": _text(row["repository"], required=True),
            "website": _text(row["website"], required=True),
            "summary": _text(row["summary"], required=True),
            "problem": _text(row["problem"], required=True),
            "ideal_for": _text(row["ideal_for"], required=True),
            "department_ids": _list(row.get("department_ids")),
            "category_ids": _list(row.get("category_ids")),
            "buyer_roles": _list(row.get("buyer_roles")),
            "company_sizes": _list(row.get("company_sizes")),
            "deployment_modes": _list(row.get("deployment_modes")),
            "service_types": _list(row.get("service_types")),
            "license_id": _text(row["license_id"], required=True),
            "maturity": _text(row["maturity"], required=True),
            "editorial_score": int(row["editorial_score"]),
            "english_support": _text(row["english_support"], required=True),
            "greek_support": _text(row["greek_support"], required=True),
            "greek_evidence": _text(row.get("greek_evidence")),
            "edition_boundary": _text(row.get("edition_boundary")),
            "featured": _bool(row.get("featured")),
            "last_verified_at": _date(row["last_verified_at"]),
            "extra_keywords_en": _list(row.get("extra_keywords_en")),
            "extra_keywords_el": _list(row.get("extra_keywords_el")),
            "negative_keywords": _list(row.get("negative_keywords")),
            "delivery_fit": _text(row["delivery_fit"], required=True),
            "active": _bool(row.get("active"), default=True),
        }
        for row in _rows(workbook, "Products")
    ]

    payload: dict[str, Any] = {
        "schema_version": schema_version,
        "source": {
            "name": dictionary.get("source_name", "opensource-for-business"),
            "snapshot_date": dictionary.get("source_snapshot_date", "2026-08-02"),
            "ownership": dictionary.get("ownership", "Opportunity Finder"),
        },
        "departments": departments,
        "categories": categories,
        "licenses": licenses,
        "repository_health": repository_health,
        "products": products,
    }
    _validate_relationships(payload)
    payload["catalog_version"] = catalog_version(payload)
    try:
        return SoftwareCatalog.model_validate(payload).model_dump(mode="json")
    except ValidationError as exc:
        raise ValueError(f"Workbook schema validation failed: {exc}") from exc


def _validate_relationships(payload: dict[str, Any]) -> None:
    for name, expected in EXPECTED_COUNTS.items():
        actual = len(payload[name])
        if actual != expected:
            raise ValueError(f"Expected {expected} {name}, found {actual}")

    def unique(items: list[dict[str, Any]], key: str, label: str) -> set[str]:
        values = [str(item[key]) for item in items]
        if len(values) != len(set(values)):
            raise ValueError(f"Duplicate {label} values detected")
        return set(values)

    department_ids = unique(payload["departments"], "id", "department id")
    category_ids = unique(payload["categories"], "id", "category id")
    license_ids = unique(payload["licenses"], "id", "license id")
    product_slugs = unique(payload["products"], "slug", "product slug")
    unique(payload["products"], "repository", "product repository URL")
    health_slugs = unique(payload["repository_health"], "slug", "repository health slug")
    if product_slugs != health_slugs:
        raise ValueError("Repository Health slugs must match Products slugs exactly")
    for category in payload["categories"]:
        if category["department_id"] not in department_ids:
            raise ValueError(f"Unknown department on category {category['id']}")
    for product in payload["products"]:
        if not product["department_ids"] or not set(product["department_ids"]).issubset(department_ids):
            raise ValueError(f"Invalid departments on product {product['slug']}")
        if not product["category_ids"] or not set(product["category_ids"]).issubset(category_ids):
            raise ValueError(f"Invalid categories on product {product['slug']}")
        if product["license_id"] not in license_ids:
            raise ValueError(f"Unknown license on product {product['slug']}")
        if product["greek_support"] in {"verified", "partial"} and not product["greek_evidence"]:
            raise ValueError(f"Greek evidence is required for product {product['slug']}")
        for field in ("repository", "website"):
            if not _valid_url(product[field]):
                raise ValueError(f"Invalid {field} URL on product {product['slug']}")
        if product["greek_evidence"] and not _valid_url(product["greek_evidence"]):
            raise ValueError(f"Invalid Greek evidence URL on product {product['slug']}")
    for license_profile in payload["licenses"]:
        if not _valid_url(license_profile["evidence"]):
            raise ValueError(f"Invalid evidence URL on license {license_profile['id']}")


def _valid_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.hostname)


def write_snapshot(workbook_path: Path, output_path: Path) -> dict[str, Any]:
    payload = workbook_payload(workbook_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate the software catalog workbook and compile runtime JSON")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check", action="store_true", help="Validate that the generated JSON is current")
    args = parser.parse_args()
    payload = workbook_payload(args.workbook)
    serialized = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    if args.check:
        if not args.output.exists() or args.output.read_text(encoding="utf-8") != serialized:
            raise SystemExit("software catalog JSON is stale; regenerate it from the workbook")
        print(f"Catalog OK: {len(payload['products'])} products, {payload['catalog_version']}")
        return 0
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(serialized, encoding="utf-8")
    print(f"Wrote {len(payload['products'])} products to {args.output} ({payload['catalog_version']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
